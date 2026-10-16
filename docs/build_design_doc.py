"""
설계서(.docx) 생성 스크립트 — 단계가 끝날 때마다 내용을 추가하고 다시 실행한다.

  .venv\\Scripts\\python docs\\build_design_doc.py        → docs/설계서.docx
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent / "설계서.docx"
FONT = "맑은 고딕"
MONO = "Consolas"
DOC_VERSION = "3단계 완료판 (2026-10-11, 파이 OS 확정 반영 2026-10-14)"


# ================================================================ 서식 도우미
def set_font(run_or_style, name, size=None):
    font = run_or_style.font
    font.name = name
    rpr = run_or_style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size:
        font.size = Pt(size)


def setup_styles(doc):
    set_font(doc.styles["Normal"], FONT, 10)
    for i, size in ((1, 16), (2, 13), (3, 11)):
        style = doc.styles[f"Heading {i}"]
        set_font(style, FONT, size)
        style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    set_font(doc.styles["Title"], FONT, 24)


def para(doc, text, bold=False, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):  # (굵은 머리말, 설명)
            p.add_run(item[0]).bold = True
            p.add_run(" — " + item[1])
        else:
            p.add_run(item)


def numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def code(doc, text):
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(12)
        set_font(p.add_run(line if line else " "), MONO, 8.5)
    doc.add_paragraph()


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_font(r, FONT, 9)
    if widths:
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Pt(w)
    doc.add_paragraph()
    return t


# ================================================================ 내용
def overview(doc):
    doc.add_heading("1. 개요", 1)
    para(doc, "목표: 바코드 인식 → 라즈베리파이4(MQTT 브로커) → 데이터베이스 및 아두이노 UNO")
    para(doc, "물류창고에서 카메라로 1D 바코드(상품 EAN-13, 위치·팔레트 Code128 등)를 인식하고, "
              "MQTT로 중앙(라즈베리파이4)에 전달해 데이터베이스에 기록하며, 아두이노 UNO로 현장 표시·알림을 한다.")
    doc.add_heading("1.1 전체 구성", 2)
    code(doc, """
[카메라] ─► 스캐너(Python, 라즈베리파이) ─► outbox(SQLite) ─► MQTT 발행 (QoS 1)
                                                                  │
                                                     Mosquitto 브로커 (라즈베리파이4)
                                                                  │
                         ┌────────────────────────────────────────┴──────────────┐
                (2단계) DB 저장 서비스 ─► 데이터베이스        (3단계) 시리얼 브리지 ─► 아두이노 UNO
""")
    doc.add_heading("1.2 단계별 진행 현황", 2)
    table(doc, ["단계", "내용", "상태"], [
        ["1단계", "바코드 인식 → MQTT 발행", "완료 (Windows 로컬 검증. 파이·카메라 실측은 미검증)"],
        ["2단계", "MQTT 구독 → DB 저장", "완료 (Windows 로컬 검증. 파이 systemd 동작은 미검증)"],
        ["3단계", "아두이노 UNO 연동", "완료 (Windows 로컬 검증: 컴파일·가짜 UNO. 실제 보드는 미검증)"],
    ], widths=[50, 170, 230])


def environment(doc):
    doc.add_heading("2. 개발·실행 환경", 1)
    table(doc, ["구분", "내용"], [
        ["Windows 작업 PC", "코드 작성, 단위·통합 테스트, 로컬 Mosquitto 테스트, 설계서 생성. "
                            "Python 3.11.9 (D:\\Tools\\Python311), venv(.venv), requirements-dev.txt"],
        ["라즈베리파이4", "실제 실행 환경. Raspberry Pi OS Bookworm 64비트(aarch64, Python 3.11)로 확정 (2026-10-14). "
                         "apt: picamera2·libzbar0·mosquitto / pip(venv --system-site-packages): requirements-pi.txt"],
        ["아두이노 UNO", "Arduino C++ 스케치. arduino-cli 1.5.1 + arduino:avr 1.8.8 + LiquidCrystal I2C 1.1.2 "
                        "(sketch.yaml 프로필). CH340/CH341 호환보드 예정. "
                        "UNO 프로토콜 코드 PC 테스트용 w64devkit g++ 16.2 (D:\\Tools\\w64devkit)"],
        ["로컬 브로커", "Mosquitto 2.1.2 (D:\\Tools\\mosquitto). Windows 서비스는 수동 시작으로 두고, "
                       "테스트가 임시 포트로 직접 실행·종료한다"],
        ["버전 관리", "Git + GitHub 비공개 저장소 JaeFree/major-project-4-2"],
    ], widths=[90, 360])
    para(doc, "버전 고정 이유: Bookworm의 apt picamera2는 numpy 1.x로 빌드되어 있어 venv에서 numpy<2를 유지해야 한다. "
              "opencv-python 4.12부터 numpy>=2가 필요하므로 OpenCV는 4.11.0.86으로 고정했다. "
              "또 Bookworm의 apt OpenCV(4.6)에는 cv2.barcode가 없어 파이에서도 pip의 opencv-python-headless를 쓴다.")


def stage1(doc):
    doc.add_heading("3. 1단계: 바코드 인식 → MQTT 발행", 1)

    doc.add_heading("3.1 구현한 기능", 2)
    bullets(doc, [
        ("바코드 인식", "OpenCV BarcodeDetector와 ZBar(pyzbar)를 매 프레임 함께 실행해 결과를 합친다. "
                       "종류 표기는 ZBar 방식(EAN13, CODE128 …)으로 통일"),
        ("중복 발행 방지", "ScanTracker — interval(계속 보이면 cooldown마다) / presence(등장마다 1회), "
                          "confirm(연속 인식 횟수)으로 오인식 차단"),
        ("오프라인 우선 전송", "스캔을 디스크 outbox(SQLite)에 먼저 기록한 뒤 발행. 브로커·네트워크가 끊겨도 "
                              "유실 없이 재연결 시 순서대로 전송"),
        ("멱등성", "메시지마다 uuid(id)와 스캐너별 일련번호(seq). 재전송으로 중복이 생겨도 받는 쪽에서 한 번만 반영"),
        ("상태 알림", "status 토픽(retain): online / offline(정상 종료) / lost(비정상 끊김, 브로커가 LWT로 발행)"),
        ("브로커 보안", "사용자 인증, 토픽 ACL(스캐너는 자기 토픽에만 쓰기), 영속성 설정"),
        ("paho-mqtt 1.x/2.x 호환", "콜백 형식 차이를 mqtt_compat 계층에서 흡수 (둘 다 테스트)"),
        ("카메라", "Picamera2 / USB 웹캠(OS별 백엔드: Linux V4L2, Windows DirectShow) / 이미지 폴더(시연·테스트)"),
        ("코드 재사용", "기존 barcode_scanner.py의 메인 루프를 ScanLoop로 분리. 기존 단독 실행판(CSV+SQLite)도 유지"),
    ])

    doc.add_heading("3.2 구성 모듈", 2)
    table(doc, ["파일", "역할"], [
        ["barcode_system/camera.py", "Camera(Picamera2·USB), ImageFolderCamera, OS별 백엔드 선택"],
        ["barcode_system/decoding.py", "프레임 검증, BarcodeReader(OpenCV+ZBar), 표기 통일(UPC-A→EAN13), 그리기"],
        ["barcode_system/tracker.py", "ScanTracker (중복 판단)"],
        ["barcode_system/scan_loop.py", "ScanLoop: 카메라→인식→중복 판단→on_scan 콜백, 카메라 자동 재연결, 미리보기"],
        ["barcode_system/messages.py", "토픽 이름, 메시지 생성·검증(JSON)"],
        ["barcode_system/outbox.py", "SQLite outbox (기록, 미전송 조회, 전송 완료 표시, 보관 기간 정리)"],
        ["barcode_system/mqtt_compat.py", "paho-mqtt 1.x/2.x 호환 계층"],
        ["barcode_system/publisher.py", "발행 스레드: 재연결, inflight 제한, PUBACK 처리, 재발행, status/LWT"],
        ["barcode_system/scanner_app.py", "실행 진입점, 설정(INI + 명령줄) 처리"],
        ["barcode_scanner.py", "기존 단독 실행판 (CSV+SQLite). 위 모듈을 공유"],
        ["config/scanner.example.ini", "설정 예시"],
        ["mosquitto/pi/barcode.conf, acl, setup_broker.sh", "파이 브로커 설정·설치 스크립트"],
        ["tests/*, scripts/test_paho1.ps1", "테스트 (3.8 참조)"],
    ], widths=[170, 280])

    doc.add_heading("3.3 동작 방식", 2)
    numbered(doc, [
        "ScanLoop가 카메라 프레임을 받아 형식을 검증하고(validate_frame) 두 엔진으로 바코드를 찾는다.",
        "ScanTracker가 '지금 발행할지' 판단한다. 시간 간격은 time.monotonic()으로 계산해 시스템 시계 변경(NTP)에 영향받지 않는다.",
        "발행 대상이면 outbox에 한 트랜잭션으로 기록한다(seq 발급 → uuid·시각을 담은 JSON 작성). "
        "WAL + synchronous=FULL이라 commit 후에는 전원이 꺼져도 남는다.",
        "발행 스레드가 outbox의 미전송 메시지를 seq 순서로 QoS 1 발행한다. 연결이 끊겨 있으면 발행하지 않고 기다린다.",
        "브로커의 PUBACK을 받으면 outbox에 sent_at을 기록한다. 응답 대기 메시지는 최대 20개(inflight)로 제한한다.",
        "30초 안에 PUBACK이 없으면 같은 메시지(같은 id)를 다시 발행한다. 재연결 후에는 paho가 스스로 재전송할 수 있어 기한을 새로 준다.",
        "연결·재연결은 paho 네트워크 스레드가 1초→최대 30초 간격으로 재시도한다. 첫 연결이 실패해도 스캔은 계속된다.",
        "정상 종료 시 남은 메시지를 최대 2초 동안 보내고 offline 상태를 남긴다. 남은 메시지는 다음 실행 때 전송된다.",
        "전송 완료된 기록은 7일(설정 가능) 뒤 정리한다. 미전송 기록은 지우지 않는다.",
    ])
    para(doc, "전달 보장: at-least-once(최소 한 번). 중복은 id로 제거한다(2단계 DB에서 id를 UNIQUE로 저장 예정).", bold=True)
    para(doc, "paho 콜백은 이벤트 큐에 넣기만 하고 실제 처리는 발행 스레드 하나가 한다. "
              "inflight 상태를 스레드 하나만 다루므로 잠금이 필요 없고, publish()가 반환되기 전에 PUBACK이 먼저 오는 경우에도 안전하다.")

    doc.add_heading("3.4 토픽과 메시지 형식", 2)
    table(doc, ["토픽", "방향", "QoS", "retain", "내용"], [
        ["warehouse/scanner/{scanner_id}/scan", "스캐너→브로커", "1", "X", "스캔 이벤트"],
        ["warehouse/scanner/{scanner_id}/status", "스캐너→브로커", "1", "O", "online / offline / lost"],
    ], widths=[170, 70, 30, 40, 140])
    para(doc, "scanner_id는 영문·숫자·_·- 1~64자이며 MQTT 사용자 이름과 같게 둔다(ACL 규칙). "
              "창고가 여러 곳이 되면 warehouse/{site}/scanner/... 로 단계를 추가할 수 있다.")
    code(doc, """
scan   {"v":1,"id":"3f2b9c1e-...","scanner_id":"dock-01","seq":1532,
        "ts":"2026-10-11T20:41:03.123+09:00","type":"EAN13","data":"8801234567893"}
status {"v":1,"scanner_id":"dock-01","state":"online","ts":"2026-10-11T20:40:00.000+09:00"}
""")
    table(doc, ["필드", "설명"], [
        ["v", "메시지 형식 버전 (현재 1)"],
        ["id", "uuid4. outbox 기록 시 생성되고 재전송해도 바뀌지 않는다 → 중복 제거 키"],
        ["seq", "스캐너별 일련번호(1부터, 재시작해도 이어짐) → 2단계에서 누락 확인"],
        ["ts", "스캔 시각(밀리초·시간대 포함). 파이는 RTC가 없어 NTP 동기화 전엔 틀릴 수 있음 → 2단계에서 수신 시각도 저장"],
        ["type / data", "바코드 종류(ZBar 표기) / 값(UTF-8, 최대 4096자)"],
    ], widths=[70, 380])

    doc.add_heading("3.5 브로커(Mosquitto) 설정", 2)
    bullets(doc, [
        ("인증", "allow_anonymous false, password_file. 사용자는 setup_broker.sh가 비밀번호를 입력받아 등록"),
        ("영속성", "persistence true(데비안 기본), max_queued_messages 100000. "
                   "저장 주기는 1단계에서 60초였고 2단계에서 '변경할 때마다'로 바꿨다(4.5 참조)"),
        ("크기 제한", "message_size_limit 64KB"),
        ("TLS", "이번 단계 미적용 (내부망 전제). 필요 시 8883 리스너로 추가"),
    ])
    table(doc, ["사용자", "권한"], [
        ["스캐너 (dock-01 등)", "warehouse/scanner/<자기 이름>/scan, /status 쓰기 (pattern %u)"],
        ["db-writer (2단계)", "warehouse/scanner/+/scan, +/status 읽기"],
        ["monitor", "warehouse/# 읽기 (점검용)"],
    ], widths=[120, 330])

    doc.add_heading("3.6 설정과 실행", 2)
    code(doc, r"""
# 설정: config/scanner.example.ini 를 config/scanner.ini 로 복사해 수정 (명령줄 옵션이 우선)
python -m barcode_system.scanner_app --config config/scanner.ini
# 웹캠 없이 시연: 이미지 폴더를 카메라처럼 사용
python -m barcode_system.scanner_app --id dock-01 --camera files --images samples --host 127.0.0.1
# 비밀번호: [mqtt] password_file 의 첫 줄 또는 환경 변수 BARCODE_MQTT_PASSWORD (명령줄에는 쓰지 않음)
# 종료: Ctrl+C, systemd stop(SIGTERM), Windows Ctrl+Break → 정상 종료(남은 메시지 전송, offline 발행)
""")

    doc.add_heading("3.7 사용 기술", 2)
    table(doc, ["기술", "용도"], [
        ["Python 3.11", "스캐너 전체"],
        ["OpenCV 4.11 (cv2.barcode)", "EAN/UPC 계열 위치 검출·해독, 영상 처리"],
        ["ZBar (pyzbar 0.1.9)", "Code128, Code39, I25, Codabar 등 1D 바코드 해독"],
        ["Picamera2 / V4L2 / DirectShow", "카메라 입력"],
        ["SQLite (WAL, synchronous=FULL)", "outbox — 전원 차단에도 기록 보존"],
        ["MQTT 3.1.1, QoS 1, retain, LWT", "메시지 전달, 상태 알림"],
        ["Eclipse Mosquitto 2.x", "브로커 (인증, ACL, 영속성)"],
        ["paho-mqtt 1.6.1 / 2.1", "MQTT 클라이언트 (두 버전 호환)"],
        ["pytest, python-barcode", "테스트, 합성 바코드 이미지 생성"],
    ], widths=[170, 280])

    doc.add_heading("3.8 테스트 결과 (Windows, 2026-10-11)", 2)
    para(doc, "전체 112개 통과 (paho-mqtt 2.1.0). MQTT 관련 29개는 paho-mqtt 1.6.1에서도 통과.", bold=True)
    table(doc, ["테스트 파일", "개수", "확인 내용"], [
        ["test_env.py", "4", "패키지, OpenCV·ZBar 인식, 로컬 브로커 왕복"],
        ["test_tracker.py", "10", "interval/presence/confirm, cooldown=0 회귀, 이력 정리"],
        ["test_decoding.py", "32", "EAN13·UPC-A·Code128·Code39, 한 화면 2개, 인식률, 프레임 검증, 좌표 정규화"],
        ["test_camera.py", "8", "OS별 백엔드 선택, 이미지 폴더 카메라"],
        ["test_messages.py", "26", "토픽, 메시지 생성·검증, 잘못된 입력 거부"],
        ["test_outbox.py", "7", "순서, 재시작 후 보존, seq 재사용 없음, 정리, 롤백, 동시 기록"],
        ["test_publisher.py", "9", "inflight 제한, PUBACK 시간 초과 재발행, 끊김 중 발행 중단, "
                                   "네트워크 단절 23건 유실 0, 브로커 재시작·브로커 없이 시작, LWT"],
        ["test_broker_security.py", "5", "익명·틀린 비밀번호 거부, ACL(사칭·읽기 전용 쓰기 차단), 인증 상태에서 LWT"],
        ["test_scanner_app.py", "11", "설정 우선순위·오류, 실제 프로세스 end-to-end, 강제 종료→재시작 유실 0, "
                                     "정상 종료, 기존 단독 실행판 동작"],
    ], widths=[120, 35, 295])
    para(doc, "주요 시나리오 결과")
    bullets(doc, [
        "네트워크 단절(TCP 프록시로 스캐너 연결만 차단) 중 20건 → 단절 동안 outbox에 20건 보관, 복구 후 23건 모두 도착(고유 id 기준).",
        "스캐너 프로세스 강제 종료 → 브로커가 lost 발행 → 재시작 시 남은 메시지 전송, seq 1..N 빈 번호 없음, 수신 누락 0.",
        "합성 EAN-13 인식률(40개): ZBar 40/40, OpenCV 단독 최대 36/40 → 두 엔진 동시 실행(auto)이 필요함을 확인.",
    ])

    doc.add_heading("3.9 발견한 오류와 수정", 2)
    table(doc, ["오류", "원인", "수정"], [
        ["auto 모드에서 UPC-A 바코드가 값이 다른 두 건으로 저장·발행됨 (기존 코드)",
         "OpenCV는 12자리 UPCA, ZBar는 앞에 0을 붙인 13자리 EAN13으로 보고 → 값이 달라 중복 제거 실패",
         "canonical(): 12자리 UPCA를 0+값, EAN13(GTIN-13)으로 통일. 회귀 테스트 추가"],
        ["presence 모드에서 cooldown=0이면 매 프레임 저장 (기존 코드)",
         "연속 프레임 사이에도 시간 간격이 있어 '0초 이상 안 보임' 조건이 항상 참",
         "'사라짐' 기준을 max(cooldown, 0.5초)로. 회귀 테스트 추가"],
        ["Windows에서 --show 미리보기가 항상 꺼짐 (기존 코드)",
         "디스플레이 확인이 DISPLAY 환경 변수(리눅스 전용)만 봄",
         "Linux에서만 DISPLAY/WAYLAND_DISPLAY 확인"],
        ["Windows에서 USB 카메라를 열 수 없음 (알려진 문제)",
         "cv2.CAP_V4L2 고정",
         "OS별 백엔드 자동 선택(--backend로 변경 가능)"],
        ["미리보기를 켜면 원본 프레임에 외곽선이 그려짐",
         "BGR 프레임은 to_bgr()가 같은 배열을 돌려줌 → 이미지 폴더 카메라처럼 프레임을 재사용하면 다음 인식에 영향",
         "그리기 전에 복사"],
        ["합성 이미지 테스트에서 OpenCV가 인식 실패",
         "OpenCV 검출기는 밝기 기울기로 위치를 찾아 경계가 완전히 날카로운 이미지를 놓침",
         "테스트 이미지에 여백·약한 흐림 적용(camera_like). 인식률 테스트로 특성 기록"],
        ["paho 1.x 테스트 스크립트(PowerShell) 실행 시 구문 오류",
         "Windows PowerShell 5.1은 BOM 없는 UTF-8을 CP949로 읽어 한글 문자열이 깨짐",
         "스크립트를 UTF-8(BOM)으로 저장"],
    ], widths=[140, 160, 150])

    doc.add_heading("3.10 미검증 항목 (하드웨어 없음)", 2)
    bullets(doc, [
        "실제 카메라(Pi 카메라 모듈, USB 웹캠)로 실물 바코드 인식 — 초점·조명·거리·움직임에 따른 인식률",
        "라즈베리파이4에서의 설치(requirements-pi.txt, apt), 실행 속도(FPS), CPU 사용률",
        "파이의 Mosquitto 설치 스크립트(setup_broker.sh)와 설정 — Windows의 Mosquitto 2.1.2로 같은 ACL만 검증",
        "실제 네트워크(Wi-Fi) 단절·전원 차단 상황 — 로컬 TCP 프록시와 프로세스 강제 종료로 흉내 냄",
        "UPC-E, I25, Codabar, DataBar — 생성기가 없어 합성 이미지 테스트를 하지 못함",
    ])

    doc.add_heading("3.11 알려진 제약과 이후 과제", 2)
    bullets(doc, [
        "MQTT 3.1.1에서는 ACL로 거부된 발행도 PUBACK을 받아 스캐너는 전송 성공으로 기록한다. "
        "사용자 이름과 scanner_id가 다르면 경고를 출력한다.",
        "ACL의 pattern 규칙은 모든 사용자에 적용되어 db-writer·monitor도 warehouse/scanner/<자기 이름>/... 에 쓸 수 있다. "
        "→ 2단계 db-writer가 토픽의 scanner_id와 메시지의 scanner_id가 다르면 저장하지 않고 rejected에 남긴다.",
        "브로커가 비정상 종료되면 마지막 자동 저장(60초) 이후의 브로커 내부 큐를 잃을 수 있었다. "
        "→ 2단계에서 '변경할 때마다 저장'으로 바꾸고 테스트로 확인했다 (4.5).",
        "시각(ts)은 파이 시계 기준이다. NTP 동기화 전 스캔은 시각이 틀릴 수 있다.",
    ])


def stage2(doc):
    doc.add_heading("4. 2단계: MQTT 구독 → DB 저장", 1)

    doc.add_heading("4.1 구현한 기능", 2)
    bullets(doc, [
        ("DB 저장 서비스(db-writer)", "모든 스캐너의 scan·status를 구독해 SQLite에 스캔 원본을 기록 (파이에서 상시 실행)"),
        ("유실 방지", "영속 세션 + QoS 1 + 'DB commit 후 PUBACK'. db-writer가 꺼져 있던 동안의 메시지도 다시 켜면 저장"),
        ("중복 제거", "scans.id(uuid)를 PRIMARY KEY로 → 재전송·중복 발행이 와도 1행"),
        ("검증·사칭 차단", "형식 오류, 알 수 없는 토픽, 토픽과 내용의 scanner_id 불일치 → rejected 테이블에 원본·사유 보관"),
        ("DB 장애 대비", "재시도 → 보조 파일(JSONL) → 주기적으로 DB 반영. 둘 다 실패하면 PUBACK 전에 종료(fail-stop)"),
        ("스캐너 상태", "scanners 테이블: online/offline/lost, 마지막 수신 시각, 마지막 seq, 건수"),
        ("관리 도구(db_tool)", "스캔 조회, 스캐너 상태, seq 누락 구간, 거부 메시지, CSV 내보내기, 삭제(미리보기·백업)"),
        ("자동 실행", "systemd 서비스(db-writer, scanner)와 설치 스크립트"),
    ])

    doc.add_heading("4.2 구성 모듈", 2)
    table(doc, ["파일", "역할"], [
        ["barcode_system/store.py", "DB 스키마, 저장(중복 제거), 상태 갱신, 거부 기록, 조회, 누락 구간, 삭제·백업"],
        ["barcode_system/spool.py", "DB 기록 실패 시 보조 파일 (fsync, 원자적 교체)"],
        ["barcode_system/subscriber.py", "DbWriter: 영속 세션 구독, 메시지 분류, 재시도·보조 파일·fail-stop"],
        ["barcode_system/db_writer_app.py", "실행 진입점, 설정"],
        ["barcode_system/db_tool.py", "관리 도구 (조회·내보내기·삭제)"],
        ["barcode_system/config.py", "스캐너·db-writer 공용 설정 처리 (명령줄 > INI > 기본값)"],
        ["config/db_writer.example.ini", "설정 예시"],
        ["deploy/systemd/*.service, deploy/install_services.sh", "부팅 시 자동 실행, 비정상 종료 시 재시작"],
        ["mosquitto/pi/barcode.conf", "영속성 저장 주기 변경, 영속 세션 만료 기간 추가"],
    ], widths=[170, 280])

    doc.add_heading("4.3 동작 방식", 2)
    numbered(doc, [
        "db-writer는 고정 client id(db-writer)와 clean_session=False로 접속해 warehouse/scanner/+/scan, "
        "+/status를 QoS 1로 구독한다. 브로커는 이 세션을 기억하고, db-writer가 꺼져 있는 동안 온 메시지를 보관한다.",
        "메시지를 받으면 토픽 모양 → JSON 형식 → 토픽과 내용의 scanner_id 일치 순서로 검증한다. 실패하면 rejected에 기록.",
        "scan은 INSERT OR IGNORE로 저장한다. 이미 같은 id가 있으면 '중복'(내용이 다르면 '충돌' 경고)으로 센다.",
        "status는 받은 순서대로 scanners에 반영한다. LWT(lost)의 ts는 연결할 때 미리 만든 값이라 실제 끊긴 시각보다 "
        "이르므로, ts 비교 대신 수신 순서와 수신 시각(state_at)을 쓴다.",
        "paho는 on_message 콜백이 끝난 뒤 PUBACK을 보낸다(1.x·2.x 동일). 콜백 안에서 commit까지 끝내므로 "
        "기록 전에 죽으면 브로커가 같은 메시지를 다시 보낸다.",
        "DB가 잠겨 있으면 0.1·0.5·1초 간격으로 재시도하고, 그래도 실패하면 보조 파일에 fsync로 남긴다. "
        "보조 파일은 10초마다 DB로 옮긴다(시작·종료 때도).",
        "보조 파일까지 실패하면 PUBACK을 보내기 전에 프로세스를 즉시 종료(종료 코드 3)하고, systemd가 3초 뒤 재시작한다.",
    ])
    para(doc, "결과적으로 스캐너 outbox → 브로커 → DB 전 구간이 at-least-once 전달 + id 기준 중복 제거 = "
              "'정확히 한 번 반영'이 된다.", bold=True)

    doc.add_heading("4.4 DB 구조 (SQLite, WAL, 스키마 버전 1)", 2)
    table(doc, ["테이블", "컬럼", "설명"], [
        ["scans", "id (PK)", "메시지 uuid — 중복 제거 키"],
        ["", "scanner_id, seq", "스캐너와 일련번호 (누락 확인)"],
        ["", "type, data", "바코드 종류(ZBar 표기)와 값"],
        ["", "scanned_at", "스캔 시각(UTC로 변환) — 정렬·기간 조회·삭제 기준"],
        ["", "scan_ts", "스캐너가 보낸 원래 시각 문자열(시간대 포함)"],
        ["", "received_at", "db-writer 수신 시각(UTC) — 파이에 RTC가 없어 스캐너 시각이 틀릴 때 대비"],
        ["scanners", "scanner_id (PK), state, state_at, last_scan_at, last_seq", "스캐너별 현재 상태"],
        ["rejected", "topic, payload, reason, received_at", "거부된 메시지 원본과 사유"],
    ], widths=[60, 170, 220])
    bullets(doc, [
        "인덱스: data(바코드 조회), scanned_at(기간), (scanner_id, seq)(누락 확인)",
        "WAL 모드라 나중에 만들 웹 페이지가 읽는 중에도 기록이 막히지 않는다 (테스트로 확인).",
        "PRAGMA user_version = 1. 기말에 업무 규칙(입고·출고, 상품 정보 등) 테이블을 추가할 때 버전을 올려 이전한다. "
        "프로그램보다 새 버전의 DB는 열지 않는다.",
        "모든 시각은 'YYYY-MM-DDTHH:MM:SS.mmm+00:00' 형식이라 문자열 비교로 정렬된다.",
    ])

    doc.add_heading("4.5 브로커 영속성 변경", 2)
    para(doc, "autosave_on_changes true + autosave_interval 1 (메시지가 바뀔 때마다 디스크 저장), "
              "persistent_client_expiration 30d 추가.")
    bullets(doc, [
        "이유: 스캐너는 브로커의 PUBACK을 받으면 전송 완료로 처리한다. db-writer가 꺼져 있는 동안 브로커에 쌓인 메시지는 "
        "브로커가 비정상 종료되면 되살릴 곳이 없다.",
        "검증: 같은 시나리오(db-writer 정지 → 10건 발행 → 브로커 강제 종료 → 재시작)에서 새 설정은 10건 모두 저장, "
        "이전 설정(60초)은 0건 — 대조 테스트로 설정 변경의 효과를 확인.",
        "비용: 저장 때마다 브로커 DB 파일 전체를 다시 쓴다. 창고 스캔량(초당 수 건 이하)에서는 부담이 작고, "
        "많아지면 autosave_interval을 키운다(예: 100). 나중에 조정 가능.",
    ])

    doc.add_heading("4.6 관리 도구와 삭제 기능", 2)
    code(doc, r"""
python -m barcode_system.db_tool scans --limit 20 [--scanner dock-01] [--data 880...] [--since 2026-10-01]
python -m barcode_system.db_tool scanners          # 스캐너별 상태·건수
python -m barcode_system.db_tool gaps              # seq 빈 번호 구간 (누락 의심)
python -m barcode_system.db_tool rejected          # 거부된 메시지
python -m barcode_system.db_tool export out.csv --since 2026-10-01     # Excel용 UTF-8(BOM)
python -m barcode_system.db_tool delete --before 2026-10-01            # 미리보기: 지울 건수만 표시
python -m barcode_system.db_tool delete --before 2026-10-01 --yes --backup data/backup.db [--vacuum]
python -m barcode_system.db_tool delete --all --yes                    # 전체 삭제
python -m barcode_system.db_tool delete --rejected --yes               # 거부 기록 삭제
""")
    bullets(doc, [
        "기록은 무기한 보관하고, 삭제는 사람이 명령할 때만 한다.",
        "삭제 조건: --before/--until, --since, --scanner, --data, --id(여러 개). 조건 없이 지우려면 --all이 필요하다.",
        "기본은 미리보기이고 --yes를 붙여야 실제로 지운다. --backup은 실행 중에도 일관된 복사본을 만든다(SQLite 백업 API).",
        "DB 파일이 없으면 오류로 끝낸다(경로를 잘못 써서 빈 DB가 생기는 것 방지). db-writer 실행 중에도 사용할 수 있다.",
        "지운 seq 구간은 gaps에서 빈 번호로 보인다. 삭제 직후 브로커가 같은 메시지를 다시 보내면(드묾) 그 행은 다시 생길 수 있다.",
    ])

    doc.add_heading("4.7 설정과 실행", 2)
    code(doc, r"""
# 파이: 브로커 사용자 등록 (스캐너, db-writer, monitor)
sudo mosquitto/pi/setup_broker.sh dock-01 db-writer monitor
# 설정: config/db_writer.example.ini → config/db_writer.ini, 비밀번호 파일 config/db-writer.pass
python -m barcode_system.db_writer_app --config config/db_writer.ini
# 부팅 시 자동 실행 (systemd)
sudo deploy/install_services.sh db-writer scanner
journalctl -u barcode-db-writer -f
""")

    doc.add_heading("4.8 사용 기술", 2)
    table(doc, ["기술", "용도"], [
        ["MQTT 영속 세션(clean_session=False), QoS 1", "구독자가 꺼져 있는 동안의 메시지 보관·재전달"],
        ["SQLite (WAL, synchronous=FULL, INSERT OR IGNORE, 윈도 함수 LAG)", "원본 기록, 중복 제거, 누락 구간 계산"],
        ["SQLite 온라인 백업 API, VACUUM", "삭제 전 백업, 파일 크기 정리"],
        ["JSON Lines + fsync + os.replace", "보조 파일 (원자적 갱신)"],
        ["systemd (Restart=always)", "자동 실행·재시작"],
        ["Mosquitto autosave_on_changes", "브로커 강제 종료에도 보관 메시지 유지"],
    ], widths=[220, 230])

    doc.add_heading("4.9 테스트 결과 (Windows, 2026-10-11)", 2)
    para(doc, "전체 148개 통과 (paho-mqtt 2.1.0, 2회 반복 실행). MQTT 관련 51개는 paho-mqtt 1.6.1에서도 통과.", bold=True)
    table(doc, ["테스트 파일", "개수", "확인 내용"], [
        ["test_store.py", "11", "저장·중복·충돌, 시각 변환, 상태 반영 순서, 누락 구간, 조건 조회, 삭제·백업, "
                               "스키마 버전, 읽는 중 기록(WAL)"],
        ["test_spool.py", "3", "기록·반영, 중간 실패 시 나머지 보존, 잘린 줄 건너뛰기"],
        ["test_db_tool.py", "10", "조회·CSV, 삭제 미리보기·--yes·백업·전체·거부 기록, 잘못된 인자, 없는 DB 미생성"],
        ["test_subscriber.py", "9", "분류·사칭 거부, 중복 집계, DB 잠금→보조 파일→반영, fail-stop, "
                                    "영속 세션(꺼진 동안 20건), fail-stop 후 재전달, 실시간 DB 잠금"],
        ["test_pipeline.py", "3", "브로커 강제 종료(새 설정 유지 / 이전 설정 유실 대조), "
                                  "실제 프로세스 2개 강제 종료·재시작 후 outbox = DB"],
        ["1단계 테스트", "112", "변경 없음 (설정 처리 공용화 후에도 모두 통과)"],
    ], widths=[110, 35, 305])
    para(doc, "주요 시나리오 결과")
    bullets(doc, [
        "스캐너 앱과 db-writer를 실제 프로세스로 실행하고 db-writer 1회, 스캐너 1회 강제 종료 → "
        "outbox 23건과 DB 23건의 id가 정확히 일치, seq 1~23 빈 번호·중복 행 없음, 거부 0건, 최종 상태 offline.",
        "db-writer와 보조 파일이 모두 기록 불가 → PUBACK 전에 멈춤 → 같은 client id로 재시작한 db-writer가 같은 메시지를 받아 저장.",
    ])

    doc.add_heading("4.10 발견한 오류와 수정", 2)
    table(doc, ["오류", "원인", "수정"], [
        ["보조 파일을 읽을 수 없으면 db-writer 시작이 실패하고, 재처리 스레드도 죽음",
         "보조 파일 길이 확인(len)이 예외 처리 밖에 있었음",
         "보조 파일 처리 전체를 예외 처리 안으로. 비상용 경로 때문에 서비스가 멈추지 않게 함"],
        ["(설계 단계) 브로커 강제 종료 시 db-writer 몫으로 보관 중이던 메시지 유실",
         "1단계 설정의 60초 자동 저장 — 대조 테스트로 실제 유실(10건 중 0건 남음) 확인",
         "autosave_on_changes true + autosave_interval 1"],
    ], widths=[150, 150, 150])

    doc.add_heading("4.11 미검증 항목", 2)
    bullets(doc, [
        "라즈베리파이에서의 systemd 서비스 설치·자동 재시작, 실제 전원 차단 후 복구",
        "SD카드에서의 브로커 저장(변경마다) 부하와 DB 쓰기 성능",
        "파이 OS의 Mosquitto 버전(Bookworm은 2.0.11)에서의 설정 호환 — Windows의 2.1.2로만 검증",
    ])

    doc.add_heading("4.12 알려진 제약", 2)
    bullets(doc, [
        "db-writer는 하나만 실행한다. 같은 client id로 두 개를 띄우면 서로 접속을 빼앗는다.",
        "max_queued_messages(100000)를 넘게 쌓이면 브로커가 이후 메시지를 버린다 — db-writer가 아주 오래 꺼져 있을 때만 해당.",
        "client id를 바꾸면 이전 세션에 보관된 메시지를 받지 못한다 (30일 뒤 브로커가 정리).",
        "스캐너 outbox를 지우고 다시 만들면 seq가 1부터 다시 시작해, 누락 구간 조회가 정확하지 않을 수 있다.",
    ])


def stage3(doc):
    doc.add_heading("5. 3단계: 아두이노 UNO 연동", 1)

    doc.add_heading("5.1 브리지 필요성 검토", 2)
    table(doc, ["방법", "장점", "단점", "판단"], [
        ["A. 파이의 시리얼 브리지 (MQTT 구독 → USB 시리얼)", "추가 부품 없음. 인증·재연결·검증은 파이가, UNO는 표시만",
         "파이 서비스 1개 추가", "채택"],
        ["B. UNO에 이더넷/Wi-Fi 실드 + 직접 MQTT", "파이와 독립", "부품 구매, SRAM 2KB로 JSON·인증·재연결이 빠듯", "비권장"],
        ["C. 스캐너 프로그램이 시리얼로 직접 전송", "단순", "스캐너와 UNO가 묶여 MQTT 구조의 의미가 없어짐", "비권장"],
    ], widths=[130, 120, 130, 50])
    para(doc, "결론: UNO에는 네트워크 기능이 없으므로 파이에서 MQTT를 받아 USB 시리얼로 넘기는 브리지(uno-bridge)가 필요하다.",
         bold=True)

    doc.add_heading("5.2 구현한 기능", 2)
    bullets(doc, [
        ("DB 저장 결과 발행", "db-writer가 DB 반영 '뒤에' warehouse/db/result로 결과(inserted/duplicate/conflict/rejected) 발행 "
                             "→ UNO 표시는 'DB에 실제로 저장됨'을 뜻한다"),
        ("uno-bridge", "결과·스캐너 상태를 구독해 시리얼 프레임으로 변환·전송. 확인 응답·재전송, 포트 자동 검색(CH340/CH341/정품), "
                       "USB 분리 후 재연결, UNO 리셋 감지 시 상태 재전송, 30초보다 오래된 결과는 표시 안 함"),
        ("UNO 스케치", "I2C LCD 16x2(주소 자동 검색) + 초록/빨강 LED + 부저. 저장·거부·스캐너 lost·NO LINK 표시. "
                      "LED·부저는 millis 타이머라 표시 중에도 수신이 멈추지 않음"),
        ("시리얼 프로토콜", "한 줄 ASCII 프레임 + XOR 체크섬, 최대 63바이트. Python·C++ 양쪽 구현"),
        ("ACL", "db-writer는 결과 토픽 쓰기, uno-bridge는 결과·상태 읽기 전용"),
        ("자동 실행", "systemd 서비스(barcode-uno-bridge, dialout 그룹)"),
    ])

    doc.add_heading("5.3 구성 파일", 2)
    table(doc, ["파일", "역할"], [
        ["arduino/uno_display/uno_display.ino", "UNO 스케치 (표시·LED·부저·확인 응답)"],
        ["arduino/uno_display/protocol.h", "프레임 수신(LineReader)·해석·생성. Arduino 기능 없이 작성 → PC에서도 빌드해 테스트"],
        ["arduino/uno_display/pins.h", "배선·부품 설정 (핀, 부저 종류, LCD 주소 후보, 통신 속도)"],
        ["arduino/uno_display/sketch.yaml", "arduino-cli 프로필: arduino:avr 1.8.8, LiquidCrystal I2C 1.1.2 고정"],
        ["arduino/uno_display/README.md", "배선표, 컴파일·업로드, 단독 시험 방법"],
        ["barcode_system/uno_protocol.py", "파이 쪽 프로토콜 구현"],
        ["barcode_system/serial_bridge.py", "UnoLink(시리얼 연결), UnoBridge(MQTT → 표시 이벤트)"],
        ["barcode_system/uno_bridge_app.py, config/uno_bridge.example.ini", "실행 진입점, 설정"],
        ["barcode_system/subscriber.py, messages.py", "결과 발행 추가, result 메시지 형식"],
        ["mosquitto/pi/acl", "uno-bridge 계정, db-writer 결과 쓰기 권한"],
        ["deploy/systemd/barcode-uno-bridge.service", "자동 실행"],
        ["docs/미검증.md", "1~3단계 미검증 항목과 확인 방법 모음"],
    ], widths=[190, 260])

    doc.add_heading("5.4 동작 방식", 2)
    code(doc, """
스캐너 → 브로커 → db-writer ─(DB commit 후 warehouse/db/result)─► 브로커 → uno-bridge ─USB 시리얼─► UNO
                     스캐너 status (online/offline/lost, retain) ──────────────────┘
""")
    numbered(doc, [
        "uno-bridge는 clean_session=True로 warehouse/db/result와 warehouse/scanner/+/status를 QoS 1로 구독한다. "
        "표시는 기록이 아니므로 꺼져 있던 동안의 결과는 받지 않는다.",
        "inserted → S 프레임, rejected → X 프레임, 상태 변화 → T 프레임으로 바꿔 큐에 넣는다. "
        "duplicate/conflict는 같은 스캔의 재전달이라 표시하지 않는다.",
        "결과의 ts(db-writer 처리 시각)가 30초보다 오래됐거나, 시리얼이 끊겨 큐에서 30초 넘게 기다린 결과는 버린다.",
        "시리얼 포트를 열면 UNO가 리셋되므로 부팅 완료(R)를 최대 4초 기다린다. R이 없으면 P에 대한 확인으로 대신한다.",
        "프레임마다 일련번호 n을 붙여 보내고 0.5초 안에 A|n이 없으면 최대 3번 다시 보낸다. 연속 실패·포트 오류는 다시 연결한다.",
        "보낼 것이 없으면 5초마다 P를 보낸다. UNO는 12초 동안 아무것도 받지 못하면 NO LINK를 표시한다.",
        "UNO가 다시 R을 보내면(리셋) 또는 재연결하면 알고 있는 스캐너 상태를 다시 보낸다.",
    ])

    doc.add_heading("5.5 시리얼 프로토콜", 2)
    code(doc, """
형식   <종류>|<필드>|...*<XOR 체크섬 16진 2자리>\\n   (줄바꿈 제외 최대 63바이트, 115200bps)
파이→UNO  S|n|<type>|<scanner>|<data>     저장됨        예) S|2|EAN13|dock-01|8801234567893*34
          X|n|<scanner>|<reason>          거부됨
          T|n|<scanner>|<state>           스캐너 상태
          P|n                             연결 확인     예) P|1*1D
UNO→파이  R|uno-display|1                 부팅 완료
          A|n                             수신 확인
          E|CK / E|FMT / E|LEN            체크섬·형식 오류, 줄이 너무 김
""")
    bullets(doc, [
        "필드에는 출력 가능한 ASCII만 쓴다. '|', '*', 한글 등은 '?'로 바꾼다 (LCD도 ASCII만 표시).",
        "줄이 길면 마지막 필드(바코드 값·사유)를 자른다. UNO는 63바이트를 넘는 줄을 통째로 버리고 E|LEN을 보낸다.",
        "UNO는 처리 전에 확인 응답을 보내 브리지의 불필요한 재전송을 막는다.",
    ])

    doc.add_heading("5.6 UNO 표시와 배선", 2)
    table(doc, ["이벤트", "LCD", "LED / 부저"], [
        ["저장됨", "1줄: 종류 스캐너 / 2줄: 바코드 값", "초록 0.3초 + 짧은 삐"],
        ["거부됨", "REJECTED 스캐너 / 사유 (5초)", "빨강 빠른 깜빡임 + 긴 삐"],
        ["스캐너 lost", "스캐너 / lost (3초)", "빨강 계속 켜짐 + 삐"],
        ["브리지 끊김", "NO LINK / check bridge", "빨강 1초 주기 깜빡임"],
    ], widths=[80, 200, 170])
    table(doc, ["부품", "UNO 핀"], [
        ["I2C LCD SDA / SCL", "A4 / A5 (VCC 5V), 주소 0x27·0x3F 자동 검색"],
        ["초록 LED / 빨강 LED", "D5 / D6 (220~330Ω 직렬)"],
        ["부저", "D8 (수동 부저 기본, 능동이면 pins.h의 BUZZER_PASSIVE 0)"],
    ], widths=[150, 300])

    doc.add_heading("5.7 사용 기술", 2)
    table(doc, ["기술", "용도"], [
        ["Arduino C++ (avr-gcc 7.3, C++11), arduino-cli 1.5.1 프로필", "UNO 스케치, 버전 고정 빌드"],
        ["LiquidCrystal I2C 1.1.2, Wire(I2C)", "LCD 표시, 주소 자동 검색"],
        ["UART 115200bps, 줄 단위 ASCII + XOR 체크섬, 확인 응답·재전송", "파이 ↔ UNO 통신"],
        ["pyserial (serial_for_url, list_ports)", "시리얼 연결, USB VID:PID로 포트 검색, socket:// 시험"],
        ["w64devkit g++ 16.2 (PC)", "UNO 프로토콜 코드(protocol.h)를 PC에서 빌드해 단위 테스트"],
        ["MQTT clean session, retain 상태", "표시용 구독 (오래된 결과 무시, 현재 상태는 즉시 수신)"],
    ], widths=[230, 220])

    doc.add_heading("5.8 테스트 결과 (Windows, 2026-10-11)", 2)
    para(doc, "전체 181개 통과 (paho-mqtt 2.1.0, 2회 반복 실행). MQTT 관련 66개는 paho-mqtt 1.6.1에서도 통과.", bold=True)
    table(doc, ["테스트 파일", "개수", "확인 내용"], [
        ["test_uno_protocol.py", "15", "프레임 생성·해석, 문자 치환, 길이 제한·자르기, 오류 거부"],
        ["test_firmware.py", "3", "protocol.h C++ 단위 테스트(g++ -Wall -Wextra -Werror), Python↔C++ 대조(무작위 프레임 300여 줄 + "
                                  "손상·CRLF·초과 길이), UNO 컴파일·메모리(Flash 9,724B/30%, SRAM 728B/35%)"],
        ["test_serial_bridge.py", "12", "R 대기·확인, 확인 누락 시 재전송, R 없는 보드, 무응답, 잡음·리셋 감지, USB 분리·재연결, "
                                        "포트 자동 검색, 결과·상태 → 프레임 변환, 오래된 결과 버림, 상태 재전송, "
                                        "스캐너→db-writer→브리지→가짜 UNO 통합"],
        ["test_subscriber.py (+1)", "10", "DB 반영 후 결과 발행, 발행 실패가 저장에 영향 없음"],
        ["test_broker_security.py (+1)", "6", "결과 토픽은 db-writer만 쓰기 가능"],
        ["test_pipeline.py (+1)", "4", "실제 프로세스 3개(스캐너·db-writer·브리지) + 가짜 UNO: UNO의 S 프레임 = DB 행"],
    ], widths=[120, 35, 295])
    para(doc, "가짜 UNO(tests/fake_uno.py): 실제 펌웨어와 같은 프로토콜로 응답하는 TCP 서버. 브리지는 pyserial의 "
              "socket:// 주소로 COM 포트 대신 접속하므로, 확인 누락·재부팅·USB 분리·잡음을 하드웨어 없이 재현했다.")

    doc.add_heading("5.9 발견한 오류와 수정", 2)
    table(doc, ["오류", "원인", "수정"], [
        ["UNO 컴파일 실패: invalid types 'int[int]' for array subscript",
         "protocol.h의 배열 이름 HEX가 Arduino.h의 매크로 HEX(=16)와 충돌. PC용 g++ 빌드만으로는 드러나지 않음",
         "HEX_DIGITS로 이름 변경. 테스트에서 PC 빌드와 UNO 컴파일을 모두 실행"],
        ["Python 프로토콜 테스트 실패", "테스트 기대값 오류 ('한글\\t'은 3글자)", "기대값 수정 (구현은 정상)"],
    ], widths=[150, 170, 130])

    doc.add_heading("5.10 미검증 항목", 2)
    para(doc, "UNO가 아직 없어(CH340 또는 CH341 호환보드 예정) 다음은 검증하지 못했다. 1~3단계 전체 목록과 확인 방법은 "
              "docs/미검증.md에 모았다.")
    bullets(doc, [
        "실제 보드 업로드, CH340/CH341 드라이버·포트 자동 검색, 포트를 열 때의 리셋(R 수신)",
        "115200bps 실제 통신 오류율 (문제가 있으면 양쪽을 57600으로)",
        "LCD 주소 검색·표시, LED·부저 배선과 부저 종류, 표시 동작 전체",
        "USB 케이블 실제 분리·재연결, 파이의 dialout 권한(systemd)",
    ])

    doc.add_heading("5.11 알려진 제약", 2)
    bullets(doc, [
        "표시는 기록이 아니다. 브리지·UNO가 꺼져 있던 동안의 저장 결과는 표시되지 않는다 (DB에는 저장됨).",
        "db-writer가 결과를 발행하기 직전에 강제 종료되면 그 스캔은 다시 전달돼도 duplicate라 UNO에 표시되지 않을 수 있다.",
        "UNO는 lost 상태 스캐너를 최대 4개까지 기억한다 (pins.h의 MAX_SCANNERS).",
        "LCD 한 줄은 16자라 긴 바코드 값은 15자 + '~'로 표시된다.",
    ])


def next_steps(doc):
    doc.add_heading("6. 다음 단계", 1)
    bullets(doc, [
        ("하드웨어 검증", "카메라·파이·UNO를 갖추면 docs/미검증.md 순서대로 확인하고 결과를 기록"),
        ("기말", "업무 규칙(입고·출고, 상품 정보) 추가, 바코드가 찍힌 물품을 관리하는 웹 페이지 (SQLite DB를 그대로 사용)"),
    ])


def build():
    doc = Document()
    setup_styles(doc)
    title = doc.add_paragraph(style="Title")
    title.add_run("물류창고형 바코드 시스템 설계서")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"전공심화 4학년 2학기 프로젝트 · {DOC_VERSION}").italic = True
    table(doc, ["판", "날짜", "변경 내용"], [
        ["1단계", "2026-10-11", "환경 구성, 바코드 인식 → MQTT 발행 구현·테스트"],
        ["2단계", "2026-10-11", "MQTT 구독 → DB 저장, 관리 도구(삭제 포함), 브로커 영속성 변경, systemd"],
        ["3단계", "2026-10-11", "UNO 연동(시리얼 브리지, 스케치, 프로토콜), DB 저장 결과 발행, 미검증 목록(docs/미검증.md)"],
    ], widths=[60, 80, 310])
    overview(doc)
    environment(doc)
    stage1(doc)
    stage2(doc)
    stage3(doc)
    next_steps(doc)
    doc.save(OUT)
    print(f"생성: {OUT}")


if __name__ == "__main__":
    build()
